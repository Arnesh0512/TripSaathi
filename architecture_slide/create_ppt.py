import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank slide layout

    # Color definitions matching TripSaathi Editorial Palette
    COLOR_BG = RGBColor(248, 245, 238)         # Warm cream
    COLOR_CARD = RGBColor(255, 255, 255)       # Pure white
    COLOR_BORDER = RGBColor(230, 223, 211)     # Hairline border
    COLOR_TERRACOTTA = RGBColor(168, 56, 39)   # Primary brand
    COLOR_TEXT = RGBColor(29, 27, 24)          # Deep charcoal
    COLOR_MUTED = RGBColor(107, 100, 91)       # Slate text
    COLOR_CHIP_BG = RGBColor(248, 244, 238)    # Chip background
    COLOR_GREEN = RGBColor(35, 107, 59)        # Sage green
    COLOR_AMBER = RGBColor(189, 91, 24)        # Warning amber
    COLOR_BLUE = RGBColor(10, 108, 155)        # TinyFish blue
    COLOR_PURPLE = RGBColor(107, 59, 167)      # ChatGPT purple

    # ==========================================
    # SLIDE 1: TECHNICAL ARCHITECTURE FLOWCHART
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Background canvas fill
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()

    # Slide Header Kicker
    kicker_box = slide1.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9.0), Inches(0.3))
    tf = kicker_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRIPSAATHI CORE ARCHITECTURE • MULTI-MODAL RECOVERY PIPELINE"
    p.font.name = "Georgia"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_TERRACOTTA

    # Slide Title
    title_box = slide1.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(10.5), Inches(0.55))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Autonomous Journey Recovery: From Ingestion to Delay Resolution"
    p.font.name = "Georgia"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT

    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(10.5), Inches(0.4))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "How TinyFish Web Agents, ChatGPT Ingestion, and Deterministic Rules Automate Connected Recovery"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED

    # Brand badge in top right
    brand_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.8), Inches(0.45), Inches(1.9), Inches(0.65))
    brand_box.fill.solid()
    brand_box.fill.fore_color.rgb = COLOR_CARD
    brand_box.line.color.rgb = COLOR_BORDER
    tf = brand_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "TripSaathi"
    p.font.name = "Georgia"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "RECOVER THE JOURNEY"
    p2.font.name = "Arial"
    p2.font.size = Pt(8)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TERRACOTTA

    # Divider line under header
    line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.55), Inches(12.13), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.fill.background()

    # Lane Configuration (6 columns)
    lane_width = Inches(1.88)
    lane_gap = Inches(0.17)
    start_x = Inches(0.6)
    lane_y = Inches(1.7)
    lane_h = Inches(5.0)

    lanes_data = [
        {
            "step": "01",
            "title": "INPUT & BOOKING",
            "cards": [
                {
                    "title": "Search Request",
                    "badge": "TinyFish",
                    "badge_color": COLOR_BLUE,
                    "desc": "User submits Origin, Dest, Dates & Budget.",
                    "chip": "TinyFish searches live fares on IndiGo & Air India within budget."
                },
                {
                    "title": "1-Click Autofill",
                    "badge": "Guided",
                    "badge_color": COLOR_TERRACOTTA,
                    "desc": "Pre-fills profile into carrier link; user checks out securely."
                },
                {
                    "title": "Or: Pre-Booked",
                    "badge": "Input",
                    "badge_color": COLOR_TERRACOTTA,
                    "desc": "Uploads existing flight PDF, PNR, or forwarded email."
                }
            ]
        },
        {
            "step": "02",
            "title": "LLM PARSER",
            "cards": [
                {
                    "title": "ChatGPT Parser",
                    "badge": "gpt-4o",
                    "badge_color": COLOR_PURPLE,
                    "desc": "Parses unstructured tickets & emails into linked timeline:",
                    "chip": "• Flight AI-804: 14:15→16:30\n• Stay: Taj (Cut-off: 20:00)\n• Transit: BOM T2 Cab\n• Event: Dinner (20:00)\n\nStores into MongoDB Atlas."
                }
            ]
        },
        {
            "step": "03",
            "title": "STATUS MONITOR",
            "cards": [
                {
                    "title": "TinyFish Monitor",
                    "badge": "Crawler",
                    "badge_color": COLOR_BLUE,
                    "desc": "Autonomous background worker polling live flight radar & gate boards."
                },
                {
                    "title": "Deviation Check",
                    "badge": "Alert",
                    "badge_color": COLOR_AMBER,
                    "desc": "Monitors deviation threshold.",
                    "chip": "DISRUPTION DETECTED!\nFlight AI-804 delayed by 2h 30m.\nLanding: 16:30 → 19:00."
                }
            ]
        },
        {
            "step": "04",
            "title": "DOWNSTREAM IMPACT",
            "cards": [
                {
                    "title": "Chain Reaction",
                    "badge": "Critical",
                    "badge_color": COLOR_AMBER,
                    "desc": "Time-series engine maps ripple effects across bookings:",
                    "chip": "• Landing: 19:00\n• Transit: 1h 30m\n• Hotel Arrival: 20:30\n\nCONFLICT: 20:30 exceeds Taj 20:00 cut-off!\nDinner missed."
                }
            ]
        },
        {
            "step": "05",
            "title": "POLICY SCOUT",
            "cards": [
                {
                    "title": "TinyFish Web Scout",
                    "badge": "Web AI",
                    "badge_color": COLOR_BLUE,
                    "desc": "Dispatches crawler agents to downstream vendor sites:",
                    "chip": "Hotel Taj Policy Scraped:\n\"Free late hold until 23:00 if notified before 20:00.\"\n\nDinner Policy:\nFree reschedule before 19:00."
                }
            ]
        },
        {
            "step": "06",
            "title": "RECOVERY DECISION",
            "cards": [
                {
                    "title": "Deterministic Rules",
                    "badge": "Verified",
                    "badge_color": COLOR_GREEN,
                    "desc": "Validates buffer: 20:30 arrival < 23:00 hold (100% Feasible).",
                    "chip": "Ranked Recovery Options:\n1. Hold Stay (₹0 fee)\n2. Rebook IndiGo (+₹2.4k)\n\nUser 1-Tap Approves Plan 1."
                }
            ]
        }
    ]

    # Render Lanes and Cards
    for i, lane in enumerate(lanes_data):
        cur_x = start_x + i * (lane_width + lane_gap)

        # Lane Header Bar
        hdr_box = slide1.shapes.add_textbox(cur_x, lane_y, lane_width, Inches(0.4))
        tf = hdr_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{lane['step']} • {lane['title']}"
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT

        # Underline under lane header
        u_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, cur_x, lane_y + Inches(0.35), lane_width, Inches(0.02))
        u_line.fill.solid()
        u_line.fill.fore_color.rgb = COLOR_TERRACOTTA
        u_line.line.fill.background()

        # Render Cards inside Lane
        card_start_y = lane_y + Inches(0.45)
        total_cards = len(lane['cards'])

        for j, card in enumerate(lane['cards']):
            # Dynamic card height depending on total cards
            card_h = Inches(4.35) if total_cards == 1 else Inches(1.3) if total_cards == 3 else Inches(1.95)
            card_y = card_start_y + j * (card_h + Inches(0.15))

            card_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_x, card_y, lane_width, card_h)
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = COLOR_CARD
            card_shape.line.color.rgb = COLOR_BORDER

            tf = card_shape.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP

            # Card Title & Badge
            p = tf.paragraphs[0]
            p.text = f"{card['title']}  [{card['badge']}]"
            p.font.name = "Georgia"
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT

            # Card Description
            p2 = tf.add_paragraph()
            p2.text = card.get('desc', '')
            p2.font.name = "Arial"
            p2.font.size = Pt(9)
            p2.font.color.rgb = COLOR_MUTED

            # Payload Chip
            if 'chip' in card:
                p3 = tf.add_paragraph()
                p3.text = "\n" + card['chip']
                p3.font.name = "Courier New"
                p3.font.size = Pt(8)
                p3.font.bold = True
                p3.font.color.rgb = COLOR_TERRACOTTA

    # Slide Footer
    ftr_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.8), Inches(12.13), Inches(0.015))
    ftr_line.fill.solid()
    ftr_line.fill.fore_color.rgb = COLOR_BORDER
    ftr_line.line.fill.background()

    ftr_box = slide1.shapes.add_textbox(Inches(0.6), Inches(6.85), Inches(12.13), Inches(0.4))
    tf = ftr_box.text_frame
    p = tf.paragraphs[0]
    p.text = "CORE PRINCIPLE: AI Investigates (TinyFish) • Rules Validate (FastAPI) • People Decide (1-Tap Approval)"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_TERRACOTTA

    p_right = tf.add_paragraph()
    p_right.text = "Stack: Python FastAPI + React Native • Agentic Crawling: TinyFish Web API • Database: MongoDB Atlas"
    p_right.font.name = "Arial"
    p_right.font.size = Pt(8.5)
    p_right.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 2: DETAILED EXECUTION MATRIX
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)

    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_BG
    bg2.line.fill.background()

    # Header
    t_box2 = slide2.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(11.0), Inches(0.8))
    tf2 = t_box2.text_frame
    p = tf2.paragraphs[0]
    p.text = "TRIPSAATHI RECOVERY MATRIX"
    p.font.name = "Georgia"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_TERRACOTTA
    p2 = tf2.add_paragraph()
    p2.text = "Step-by-Step Technical Execution & Integration Matrix"
    p2.font.name = "Georgia"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT

    # Table of steps
    rows = 7
    cols = 4
    left = Inches(0.6)
    top = Inches(1.6)
    width = Inches(12.13)
    height = Inches(4.8)

    table_shape = slide2.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(4.6)
    table.columns[3].width = Inches(3.23)

    headers = ["Stage", "Primary Component", "Technical Workflow & Data Flow", "Output & System State"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TERRACOTTA
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

    matrix_rows = [
        ("01. Booking / Input", "TinyFish + DeepLinks", "User searches or inputs PNR/ticket. TinyFish crawls portals to pre-fill passenger profile directly into booking forms.", "Booking confirmed; PNR and receipt issued."),
        ("02. Extraction", "ChatGPT API (gpt-4o)", "LLM ingests unformatted ticket PDFs and booking emails. Extracts multi-modal chronology (flight times, stay cutoff, transfers).", "Normalized JSON itinerary written to MongoDB Atlas."),
        ("03. Monitoring", "TinyFish Browser Agent", "Periodic autonomous crawler queries airline flight radar and gate schedules for active status deviations.", "Status: ON_TIME or DELAYED with exact minutes."),
        ("04. Impact Engine", "Deterministic FastAPI", "Time-series calculator evaluates downstream timeline: Flight Arrival + Transfer Buffer vs. Hotel Check-in Cut-off.", "Flags critical conflicts (e.g. 20:30 arrival > 20:00 cutoff)."),
        ("05. Policy Scouting", "TinyFish Web Scout", "Dispatches crawler agents to hotel desk and activity sites to parse late-arrival policies and free cancellation terms.", "Extracted rule: \"Hold room until 23:00 at ₹0 fee\"."),
        ("06. User Approval", "React Native Client", "FastAPI calculates ranked feasible recovery plans. User receives alert and approves Option 1 with a single tap.", "Hold notice submitted; trip confirmed recovered.")
    ]

    for r, row_data in enumerate(matrix_rows, start=1):
        for c, text in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD if r % 2 == 0 else COLOR_CHIP_BG
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT
            if c == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_TERRACOTTA

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_file = r"C:\Users\Arnesh\Desktop\TripSaathi_Technical_Architecture.pptx"
    build_presentation(out_file)
